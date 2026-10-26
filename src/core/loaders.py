# -*- coding: utf-8 -*-
"""Загрузка документов для RAG-демо."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List
import warnings


@dataclass(frozen=True)
class LoadedDocument:
    id: str
    title: str
    text: str
    source: str
    document_type: str
    metadata: Dict[str, Any] = field(default_factory=dict)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Страница {number}]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                paragraphs.append(" | ".join(values))
    return "\n".join(paragraphs)


def load_document(path: str | Path) -> LoadedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    header = file_path.read_bytes()[:4]
    format_warning = None
    if suffix == ".pdf":
        if header.startswith(b"%PDF"):
            text = _read_pdf(file_path)
        else:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
            format_warning = "Файл с расширением PDF содержит обычный текст"
    elif suffix == ".docx":
        if header.startswith(b"PK"):
            text = _read_docx(file_path)
        else:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
            format_warning = "Файл с расширением DOCX содержит обычный текст"
    elif suffix in {".txt", ".md"}:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Неподдерживаемый формат документа: {suffix}")
    stat = file_path.stat()
    return LoadedDocument(
        id=file_path.stem,
        title=file_path.stem.replace("_", " ").strip().title(),
        text=text.strip(),
        source=file_path.name,
        document_type=suffix.lstrip("."),
        metadata={
            "source": file_path.name,
            "path": str(file_path.resolve()),
            "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "size_bytes": stat.st_size,
            **({"format_warning": format_warning} if format_warning else {}),
        },
    )


def load_hr_documents(folder: str | Path) -> List[LoadedDocument]:
    docs_path = Path(folder)
    if not docs_path.exists():
        return []
    if not docs_path.is_dir():
        raise ValueError(f"Путь {docs_path} не является папкой")
    documents: List[LoadedDocument] = []
    for path in sorted(docs_path.iterdir()):
        if path.is_file() and path.suffix.lower() in {".txt", ".md", ".pdf", ".docx"}:
            try:
                document = load_document(path)
            except Exception as exc:
                warnings.warn(f"Документ {path.name} пропущен: {exc}", RuntimeWarning)
                continue
            if document.text:
                documents.append(document)
    return documents


def load_documents_from_folder(folder: str | Path) -> Dict[str, str]:
    """Загружает текстовые документы из папки и возвращает словарь имя->содержимое."""
    docs_path = Path(folder)
    if not docs_path.exists():
        return {}
    if not docs_path.is_dir():
        raise ValueError(f"Путь {docs_path} не является папкой")

    try:
        return {document.source: document.text for document in load_hr_documents(docs_path)}
    except Exception as exc:
        raise RuntimeError(f"Не удалось прочитать документы из {docs_path}: {exc}") from exc
