# -*- coding: utf-8 -*-
"""Тесты для RAG-пайплайна, чанкинга и поиска."""

from typing import List

import pytest

from src.core.chunking import Chunk, chunk_text
from src.core.loaders import LoadedDocument, load_document, load_hr_documents
from src.core.rag_pipeline import RagPipeline
from src.core.retrieval import TfidfRetriever


@pytest.fixture
def sample_chunks() -> List[Chunk]:
    """Возвращает список чанков для тестов."""
    return [
        Chunk(doc_id="doc1", chunk_id=0, text="RAG помогает находить нужные фрагменты документов."),
        Chunk(doc_id="doc2", chunk_id=1, text="Для обучения достаточно небольшого набора примеров."),
    ]


def test_chunk_text_creates_expected_chunks() -> None:
    """Функция разбиения должна создавать чанки с корректными параметрами."""
    chunks = chunk_text("doc", "один два три четыре пять шесть", chunk_size_words=3, overlap_words=1)

    assert len(chunks) == 3
    assert chunks[0].chunk_id == 0
    assert chunks[0].text.split()[0] == "один"


def test_chunk_text_rejects_invalid_parameters() -> None:
    """Неверные параметры разбиения должны приводить к ошибке."""
    with pytest.raises(ValueError):
        chunk_text("doc", "текст", chunk_size_words=0, overlap_words=0)

    with pytest.raises(ValueError):
        chunk_text("doc", "текст", chunk_size_words=5, overlap_words=5)


def test_tfidf_retriever_returns_hits(sample_chunks: List[Chunk]) -> None:
    """Поисковик должен вернуть связанные чанки для запроса."""
    retriever = TfidfRetriever()
    retriever.index(sample_chunks)

    hits = retriever.search("rag", top_k=1)

    assert len(hits) == 1
    assert hits[0].chunk.doc_id == "doc1"
    assert hits[0].score >= 0.0


def test_rag_pipeline_builds_index_and_answers() -> None:
    """Пайплайн должен строить индекс и возвращать результат с чанками."""
    rag = RagPipeline(backend="tfidf")
    docs = {
        "doc1": "RAG помогает находить нужные фрагменты документов перед генерацией ответа.",
        "doc2": "Для обучения достаточно небольшого набора примеров и простой логики поиска.",
    }

    rag.build_index(docs, chunk_size_words=20, overlap_words=5)
    result = rag.ask("что такое rag", top_k=2)

    assert result.hits
    assert result.query == "что такое rag"
    assert "Вопрос:" in result.as_text()


def test_loader_reads_txt_and_docx_with_metadata(tmp_path) -> None:
    from docx import Document

    (tmp_path / "policy.txt").write_text("Правила проведения интервью", encoding="utf-8")
    document = Document()
    document.add_paragraph("Регламент адаптации сотрудника")
    document.save(tmp_path / "onboarding.docx")

    loaded = load_hr_documents(tmp_path)

    assert {item.source for item in loaded} == {"policy.txt", "onboarding.docx"}
    assert all(item.metadata["size_bytes"] > 0 for item in loaded)
    assert "Регламент адаптации" in load_document(tmp_path / "onboarding.docx").text


def test_rag_upserts_and_deletes_documents() -> None:
    rag = RagPipeline(backend="tfidf")
    first = LoadedDocument("one", "Первый", "правила найма", "one.txt", "txt")
    second = LoadedDocument("two", "Второй", "проведение интервью", "two.txt", "txt")

    rag.build_index([first])
    rag.upsert_document(second)
    assert rag.ask("интервью", top_k=1).hits[0].chunk.doc_id == "two"
    assert rag.delete_document("two") is True
    assert all(hit.chunk.doc_id != "two" for hit in rag.ask("интервью", top_k=3).hits)
    assert rag.delete_document("missing") is False


def test_rag_source_metadata_is_attached_to_chunks() -> None:
    rag = RagPipeline(backend="tfidf")
    document = LoadedDocument(
        "policy", "Политика найма", "интервью и оффер", "policy.pdf", "pdf", {"modified_at": "2026-01-01"}
    )
    rag.build_index([document])
    hit = rag.ask("оффер", top_k=1).hits[0]
    assert hit.chunk.metadata["title"] == "Политика найма"
    assert hit.chunk.metadata["source"] == "policy.pdf"
